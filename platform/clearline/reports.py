"""The client report as Word and Excel files, built from the sections the analyst chose and the reviewer approved."""

from __future__ import annotations

import io
from datetime import date

import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

INK = '#16202b'
TONES = {'accent': '#2f6fd6', 'muted': '#9aa5b1', 'green': '#2e9e6a', 'amber': '#d39b1f', 'orange': '#e2702d', 'red': '#d0433a'}
LEVELS = {'high': 'High confidence', 'medium': 'Medium confidence', 'low': 'Low confidence: a first read'}


def _short(v: float) -> str:
    a = abs(v)
    if a >= 1e9:
        return f'{v / 1e9:.1f}B'
    if a >= 1e6:
        return f'{v / 1e6:.1f}M'
    if a >= 1e4:
        return f'{v / 1e3:.0f}K'
    return f'{v:,.0f}' if a >= 10 or v == int(v) else f'{v:.1f}'


def _axis(fmt: str | None):
    return (lambda v, _: f'{v:.0f}%') if fmt == 'pct' else (lambda v, _: _short(v))


# ---------- Charts as pictures ----------


def chart_png(c: dict, currency: str) -> bytes | None:
    kind = c.get('type')
    if kind not in ('line', 'bars', 'heatmap'):
        return None
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 8.5, 'axes.edgecolor': '#c9d1da', 'axes.labelcolor': INK, 'xtick.color': '#5b6672', 'ytick.color': '#5b6672'})
    fig, ax = plt.subplots(figsize=(6.6, 2.8), dpi=170)
    ax.spines[['top', 'right']].set_visible(False)
    fmt = c.get('format')
    if kind == 'line':
        x = list(range(len(c['x'])))
        if c.get('band'):
            b = c['band']
            ax.fill_between(x, [v if v is not None else float('nan') for v in b['lo']], [v if v is not None else float('nan') for v in b['hi']], color=TONES['accent'], alpha=0.14, lw=0, label=b.get('name'))
        for s in c['series']:
            ax.plot(x, [v if v is not None else float('nan') for v in s['values']], color=TONES.get(s.get('tone') or 'accent', TONES['accent']), lw=1.8, ls='--' if s.get('dashed') else '-', label=s['name'])
        for m in c.get('marks') or []:
            ax.axvline(m['index'], color=TONES.get(m.get('tone') or 'muted', '#9aa5b1'), lw=0.8, ls=':')
            ax.annotate(m['label'], (m['index'], 1), xycoords=('data', 'axes fraction'), fontsize=7, color='#5b6672', ha='left', va='top', xytext=(3, -2), textcoords='offset points')
        step = max(1, len(x) // 8)
        ax.set_xticks(x[::step], [c['x'][i] for i in x[::step]])
        if c.get('zero'):
            ax.set_ylim(bottom=min(0, ax.get_ylim()[0]))
        ax.yaxis.set_major_formatter(_axis(fmt))
        ax.grid(axis='y', color='#e6eaee', lw=0.6)
        if len(c['series']) > 1 or c.get('band'):
            ax.legend(frameon=False, fontsize=7.5, loc='upper left')
    elif kind == 'bars':
        labels = [lab[:28] for lab in c['labels']]
        tones = c.get('tones') or []
        colors = [TONES.get(tones[i] if i < len(tones) and tones[i] else 'accent') for i in range(len(labels))]
        vals = [v if v is not None else float('nan') for v in c['values']]
        if c.get('extra'):
            c = {**c, 'extra': {**c['extra'], 'values': [v if v is not None else float('nan') for v in c['extra']['values']]}}
        if c.get('horizontal'):
            h = max(2.2, 0.28 * len(labels) + 0.6)
            fig.set_size_inches(6.6, min(h, 8))
            y = list(range(len(labels)))[::-1]
            ax.barh(y, vals, color=colors, height=0.62)
            if c.get('extra'):
                ax.barh([v - 0.3 for v in y], c['extra']['values'], color=TONES['muted'], height=0.3, label=c['extra']['name'])
            ax.set_yticks(y, labels)
            ax.xaxis.set_major_formatter(_axis(fmt))
            ax.grid(axis='x', color='#e6eaee', lw=0.6)
            if c.get('ref'):
                ax.axvline(c['ref']['value'], color=INK, lw=0.8, ls='--', label=c['ref']['label'])
        else:
            x = list(range(len(labels)))
            width = 0.4 if c.get('extra') else 0.66
            ax.bar([i - (width / 2 if c.get('extra') else 0) for i in x], vals, color=colors, width=width)
            if c.get('extra'):
                ax.bar([i + width / 2 for i in x], c['extra']['values'], color=TONES['muted'], width=width, label=c['extra']['name'])
            ax.set_xticks(x, labels, rotation=30 if len(labels) > 6 else 0, ha='right' if len(labels) > 6 else 'center')
            ax.yaxis.set_major_formatter(_axis(fmt))
            ax.grid(axis='y', color='#e6eaee', lw=0.6)
            if c.get('ref'):
                ax.axhline(c['ref']['value'], color=INK, lw=0.8, ls='--', label=c['ref']['label'])
        if c.get('ref') or c.get('extra'):
            ax.legend(frameon=False, fontsize=7.5)
    else:
        vals = [[v if v is not None else float('nan') for v in row] for row in c['values']]
        fig.set_size_inches(6.6, max(2.2, 0.3 * len(c['rows']) + 0.8))
        ax.imshow(vals, aspect='auto', cmap='Blues')
        ax.set_xticks(range(len(c['cols'])), c['cols'])
        ax.set_yticks(range(len(c['rows'])), c['rows'])
        ax.spines[:].set_visible(False)
        if len(c['rows']) * len(c['cols']) <= 220:
            top = max((v for row in vals for v in row if v == v), default=1) or 1
            for i, row in enumerate(vals):
                for j, v in enumerate(row):
                    if v == v and v:
                        ax.text(j, i, f'{v:.0f}%' if fmt == 'pct' else _short(v), ha='center', va='center', fontsize=6.5, color='white' if v > top * 0.6 else INK)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png')
    plt.close(fig)
    return buf.getvalue()


# ---------- Word ----------


def _shade(cell, hex_fill: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc.append(shd)


def _table(doc, columns: list[str], rows: list[list[str]], numeric: list[bool] | None = None, limit: int = 40) -> None:
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = 'Table Grid'
    for i, col in enumerate(columns):
        cell = t.rows[0].cells[i]
        cell.text = str(col)
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, 'EEF2F6')
    for r in rows[:limit]:
        cells = t.add_row().cells
        for i, v in enumerate(r[: len(columns)]):
            cells[i].text = str(v)
            if numeric and i < len(numeric) and numeric[i]:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
    if len(rows) > limit:
        doc.add_paragraph(f'The first {limit} of {len(rows)} rows are shown; the Excel file has them all.').runs[0].italic = True


def to_docx(rep: dict) -> bytes:
    c = rep['client']
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
        s.footer.paragraphs[0].text = f'Clearline Analytics · Confidential · prepared for {c["name"]}'
        s.footer.paragraphs[0].runs[0].font.size = Pt(8)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)

    title = doc.add_heading(c['name'], 0)
    title.runs[0].font.color.rgb = RGBColor(0x16, 0x20, 0x2B)
    meta = doc.add_paragraph()
    meta.add_run(f'Findings report · {date.today():%d %B %Y}').bold = True
    who = [f'Analyst: {c["analyst"]["name"]}' if c.get('analyst') else None, f'Reviewed by: {c["reviewer"]["name"]}' if c.get('reviewer') else None]
    if any(who):
        meta.add_run('\n' + ' · '.join(w for w in who if w))
    if c.get('question'):
        doc.add_paragraph().add_run(f'Your question: {c["question"]}').italic = True

    if rep.get('summary'):
        doc.add_heading('Summary', 1)
        for para in rep['summary'].split('\n\n'):
            doc.add_paragraph(para.strip())

    doc.add_heading('What this report covers', 1)
    for s in rep['sections']:
        doc.add_paragraph(f'{s["service"]["name"]} ({LEVELS.get(s["run"]["level"], s["run"]["level"]).split(":")[0].lower()})', style='List Bullet')

    cur = None
    for s in rep['sections']:
        res = s['run']['result']
        cur = res.get('currency', cur)
        doc.add_page_break()
        doc.add_heading(s['service']['name'], 1)
        lv = doc.add_paragraph()
        r = lv.add_run(LEVELS.get(s['run']['level'], ''))
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x5B, 0x66, 0x72)
        if s.get('note'):
            doc.add_paragraph(s['note'])
        if res.get('kpis'):
            k = res['kpis']
            t = doc.add_table(rows=2, cols=len(k))
            for i, x in enumerate(k):
                t.rows[0].cells[i].text = x['label']
                t.rows[0].cells[i].paragraphs[0].runs[0].font.size = Pt(8)
                p = t.rows[1].cells[i].paragraphs[0]
                run = p.add_run(x['value'])
                run.bold = True
                run.font.size = Pt(13)
                if x.get('note'):
                    n = p.add_run('\n' + x['note'])
                    n.font.size = Pt(7.5)
                    n.font.color.rgb = RGBColor(0x5B, 0x66, 0x72)
                for cell in (t.rows[0].cells[i], t.rows[1].cells[i]):
                    _shade(cell, 'F4F6F9')
        if res.get('findings'):
            doc.add_heading('What we found', 2)
            for f in res['findings']:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f['text'])
                if f.get('action'):
                    a = p.add_run(f'\nWhat to do: {f["action"]}')
                    a.italic = True
        for ch in res.get('charts', []):
            spec = ch['chart']
            head = doc.add_paragraph()
            head.add_run(ch['title']).bold = True
            head.paragraph_format.keep_with_next = True
            if spec.get('type') == 'table':
                _table(doc, spec['columns'], spec['rows'], spec.get('numeric'))
            else:
                try:
                    png = chart_png(spec, cur or 'EGP')
                except Exception:  # noqa: BLE001 - one chart must never stop the whole report
                    plt.close('all')
                    png = None
                    doc.add_paragraph('(This chart could not be drawn; its figures are in the Excel file.)').runs[0].italic = True
                if png:
                    doc.add_picture(io.BytesIO(png), width=Cm(16.4))
        if res.get('accuracy'):
            doc.add_paragraph().add_run(f'How sure we are: {res["accuracy"]}').italic = True
        doc.add_heading('How we worked it out', 3)
        doc.add_paragraph(s['service'].get('method', ''))
        for n in res.get('notes', []):
            p = doc.add_paragraph(n, style='List Bullet')
            p.runs[0].font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading('About your data', 1)
    doc.add_paragraph('Your files were kept encrypted on our systems and are deleted 30 days after this report is delivered. Figures are estimates from the data you sent; where the data was thin, the confidence level says so.')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Excel ----------

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='16202B')
BOLD = Font(bold=True)


def _sheet_name(name: str, used: set[str]) -> str:
    base = ''.join(ch for ch in name if ch not in '[]:*?/\\')[:28].strip() or 'Sheet'
    out, n = base, 2
    while out.lower() in used:
        out = f'{base[:25]} {n}'
        n += 1
    used.add(out.lower())
    return out


def _write_table(ws, row: int, columns: list, rows: list[list]) -> int:
    for j, col in enumerate(columns, 1):
        cell = ws.cell(row, j, str(col))
        cell.font, cell.fill = HEAD, HEAD_FILL
    for r in rows:
        row += 1
        for j, v in enumerate(r, 1):
            ws.cell(row, j, _num(v))
    return row + 2


def _num(v):
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v)
    t = s.replace(',', '').replace('−', '-').strip()
    try:
        return float(t) if t and t[0] in '-0123456789.' and not any(ch.isalpha() for ch in t) and '%' not in t else s
    except ValueError:
        return s


def chart_rows(c: dict) -> tuple[list, list[list]]:
    kind = c.get('type')
    if kind == 'line':
        cols = ['Period', *[s['name'] for s in c['series']]]
        if c.get('band'):
            cols += [f'{c["band"].get("name") or "Range"} (low)', f'{c["band"].get("name") or "Range"} (high)']
        rows = []
        for i, x in enumerate(c['x']):
            r = [x, *[s['values'][i] if i < len(s['values']) else None for s in c['series']]]
            if c.get('band'):
                r += [c['band']['lo'][i], c['band']['hi'][i]]
            rows.append(r)
        return cols, rows
    if kind == 'bars':
        cols = ['', 'Value'] + ([c['extra']['name']] if c.get('extra') else [])
        return cols, [[lab, v, *([c['extra']['values'][i]] if c.get('extra') else [])] for i, (lab, v) in enumerate(zip(c['labels'], c['values']))]
    if kind == 'table':
        return c['columns'], c['rows']
    if kind == 'heatmap':
        return ['', *c['cols']], [[r, *vals] for r, vals in zip(c['rows'], c['values'])]
    return [], []


def to_xlsx(rep: dict) -> bytes:
    c = rep['client']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = c['name']
    ws['A1'].font = Font(bold=True, size=16)
    ws['A2'] = f'Findings report · {date.today():%d %B %Y}'
    row = 4
    if rep.get('summary'):
        ws.cell(row, 1, rep['summary']).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.row_dimensions[row].height = min(400, 15 * (rep['summary'].count('\n') + len(rep['summary']) // 110 + 1))
        row += 2
    for j, h in enumerate(['Section', 'Confidence', 'Measure', 'Value', 'Note'], 1):
        cell = ws.cell(row, j, h)
        cell.font, cell.fill = HEAD, HEAD_FILL
    for s in rep['sections']:
        for k in s['run']['result'].get('kpis', []):
            row += 1
            for j, v in enumerate([s['service']['name'], s['run']['level'], k['label'], k['value'], k.get('note', '')], 1):
                ws.cell(row, j, v)
    for col, w in zip('ABCDE', (34, 12, 30, 18, 46)):
        ws.column_dimensions[col].width = w

    used = {'summary'}
    for s in rep['sections']:
        res = s['run']['result']
        sh = wb.create_sheet(_sheet_name(s['service']['name'], used))
        sh['A1'] = s['service']['name']
        sh['A1'].font = Font(bold=True, size=14)
        sh['A2'] = LEVELS.get(s['run']['level'], '')
        r = 4
        if s.get('note'):
            sh.cell(r, 1, s['note']).alignment = Alignment(wrap_text=True)
            r += 2
        if res.get('findings'):
            sh.cell(r, 1, 'What we found').font = BOLD
            r += 1
            for f in res['findings']:
                sh.cell(r, 1, f['text'])
                sh.cell(r, 2, f.get('action', ''))
                r += 1
            r += 1
        for ch in res.get('charts', []):
            cols, rows = chart_rows(ch['chart'])
            if not cols:
                continue
            sh.cell(r, 1, ch['title']).font = BOLD
            r = _write_table(sh, r + 1, cols, rows)
        if res.get('accuracy'):
            sh.cell(r, 1, f'How sure we are: {res["accuracy"]}')
        for j in range(1, 12):
            sh.column_dimensions[get_column_letter(j)].width = 16
        if res.get('findings'):
            sh.column_dimensions['A'].width, sh.column_dimensions['B'].width = 70, 50
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
