"""Made-up clients through the whole app, as an analyst would use it: upload, check columns, run every service,
add the results to the report, review it, and export Word and Excel.

    python -m tests.stress.api_run --clients 40
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import tempfile
import time
import warnings
from collections import Counter
from pathlib import Path

import pyotp

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
warnings.filterwarnings('ignore')

H = {'X-Clearline': '1'}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--clients', type=int, default=20)
    ap.add_argument('--start', type=int, default=5000)
    a = ap.parse_args()

    data = Path(tempfile.mkdtemp(prefix='clearline-stress-'))
    from clearline import config, security

    config.DATA_DIR, config.FILES_DIR, config.DB_PATH, config.MASTER_KEY_PATH = data, data / 'files', data / 'clearline.db', data / 'master.key'
    security._master = None
    from docx import Document
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from clearline.app import create_app
    from tests.stress.gen import make_client

    app = create_app()
    problems: list[tuple[int, str]] = []
    counts = Counter()

    def sign_in(email, pw):
        c = TestClient(app)
        r = c.post('/api/login', json={'email': email, 'password': pw}, headers=H).json()
        secret = r.get('secret') or SECRETS[email]
        SECRETS[email] = secret
        c.post('/api/login/code', json={'pending': r['pending'], 'code': pyotp.TOTP(secret).now()}, headers=H)
        return c

    SECRETS: dict[str, str] = {}
    boot = TestClient(app)
    boot.post('/api/setup', json={'name': 'Admin', 'email': 'admin@x.test', 'password': 'Admin-pass-2026'}, headers=H)
    admin = sign_in('admin@x.test', 'Admin-pass-2026')
    ana = admin.post('/api/users', json={'name': 'Ana Lyst', 'email': 'ana@x.test', 'role': 'analyst'}, headers=H).json()
    rev = admin.post('/api/users', json={'name': 'Rev Iewer', 'email': 'rev@x.test', 'role': 'reviewer'}, headers=H).json()
    analyst = sign_in('ana@x.test', ana['temp_password'])
    analyst.post('/api/me/password', json={'current': ana['temp_password'], 'new': 'Analyst-pass-1'}, headers=H)
    reviewer = sign_in('rev@x.test', rev['temp_password'])
    KINDS = {k['key']: k for k in admin.get('/api/kinds').json()}

    t0 = time.time()
    for seed in range(a.start, a.start + a.clients):
        cl = make_client(seed)
        p = cl.profile
        r = admin.post('/api/clients', json={'name': f'Client {seed}', 'currency': p['currency'], 'sector': 'retail', 'analyst_id': ana['id'], 'reviewer_id': rev['id']}, headers=H)
        cid = r.json()['id']
        if not cl.files:
            continue
        up = analyst.post(f'/api/clients/{cid}/files', files=[('files', (f.name, f.data, 'application/octet-stream')) for f in cl.files], headers=H)
        if up.status_code != 200:
            problems.append((seed, f'upload {up.status_code}: {up.text[:150]}'))
            continue
        for f, meta in zip(cl.files, up.json()):
            kind = f.kind
            detail = analyst.get(f'/api/files/{meta["id"]}' + ('' if meta['kind'] == kind else f'?kind={kind}'))
            if detail.status_code != 200:
                problems.append((seed, f'file detail {detail.status_code}: {detail.text[:120]}'))
                continue
            # The analyst confirms the columns; where the guess was off they pick the right one.
            fields = {x['key'] for x in KINDS[kind]['fields']}
            mapping = {**detail.json()['mapping'], **{k: v for k, v in f.headers.items() if k in fields}}
            m = analyst.put(f'/api/files/{meta["id"]}/mapping', json={'kind': kind, 'mapping': mapping, 'options': {}}, headers=H)
            counts['files'] += 1
            if m.status_code != 200 and 'No usable rows' not in m.text:
                problems.append((seed, f'mapping {kind} {m.status_code}: {m.text[:150]}'))
        svcs = analyst.get(f'/api/clients/{cid}/services').json()['services']
        added = 0
        for s in svcs:
            if s['level'] == 'none':
                continue
            rr = analyst.post(f'/api/clients/{cid}/runs', json={'service_id': s['id'], 'params': s['params']}, headers=H)
            counts['runs'] += 1
            if rr.status_code != 200:
                problems.append((seed, f'run {s["id"]} {rr.status_code}: {rr.text[:150]}'))
                continue
            body = rr.json()
            if 'error' in body:
                problems.append((seed, f'run {s["id"]} offered but failed: {body["error"][:120]}'))
                continue
            sec = analyst.post(f'/api/clients/{cid}/report/sections', json={'run_id': body['id']}, headers=H)
            if sec.status_code == 200:
                added += 1
        if not added:
            continue
        rep = analyst.get(f'/api/clients/{cid}/report').json()
        analyst.patch(f'/api/sections/{rep["sections"][0]["id"]}', json={'note': 'Checked – ملاحظة'}, headers=H)
        analyst.put(f'/api/clients/{cid}/report/summary', json={'text': 'Summary line one.\n\nملخص'}, headers=H)
        for fmt in ('docx', 'xlsx'):
            ex = analyst.get(f'/api/clients/{cid}/report.{fmt}')
            counts[fmt] += 1
            if ex.status_code != 200:
                problems.append((seed, f'export {fmt} {ex.status_code}: {ex.text[:200]}'))
                continue
            try:
                if fmt == 'docx':
                    doc = Document(io.BytesIO(ex.content))
                    text = '\n'.join(pp.text for pp in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
                else:
                    wb = load_workbook(io.BytesIO(ex.content))
                    text = '\n'.join(str(v) for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v is not None)
                bad = re.search(r'\bnan\b|\bNaT\b|\binf\b', text)
                if bad:
                    problems.append((seed, f'{fmt} contains “{bad.group(0)}”: …{text[max(0, bad.start() - 60) : bad.end() + 30]}…'))
            except Exception as e:  # noqa: BLE001
                problems.append((seed, f'{fmt} unreadable: {e}'))
        sub = analyst.post(f'/api/clients/{cid}/report/submit', headers=H)
        ok = reviewer.post(f'/api/clients/{cid}/report/review', json={'approve': True}, headers=H)
        dl = analyst.post(f'/api/clients/{cid}/report/deliver', headers=H)
        counts['reports'] += 1
        if not (sub.status_code == ok.status_code == dl.status_code == 200):
            problems.append((seed, f'workflow {sub.status_code}/{ok.status_code}/{dl.status_code}'))
        print(f'  client {seed}: {added} sections, {time.time() - t0:.0f}s', flush=True)
    log = admin.get('/api/activity').json()
    if not log['chain_intact']:
        problems.append((0, 'activity log chain broken'))
    print(f'\n{a.clients} clients · {dict(counts)} · {time.time() - t0:.0f}s · {len(problems)} problems')
    for seed, msg in problems:
        print(f'  [{seed}] {msg}')
    (Path(__file__).parent / 'out').mkdir(exist_ok=True)
    (Path(__file__).parent / 'out' / 'api_problems.json').write_text(json.dumps(problems, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
