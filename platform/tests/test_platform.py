"""End to end through the API: setup, two-step sign-in, roles, uploads, every service, the report and deletion."""

import io
import os
import time
from pathlib import Path

import pyotp
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

SAMPLES = Path(__file__).resolve().parents[1] / 'samples'
H = {'X-Clearline': '1'}


@pytest.fixture(scope='module')
def app(tmp_path_factory, monkeypatch_module):
    data = tmp_path_factory.mktemp('data')
    from clearline import config, security

    monkeypatch_module.setattr(config, 'DATA_DIR', data)
    monkeypatch_module.setattr(config, 'FILES_DIR', data / 'files')
    monkeypatch_module.setattr(config, 'DB_PATH', data / 'clearline.db')
    monkeypatch_module.setattr(config, 'MASTER_KEY_PATH', data / 'master.key')
    monkeypatch_module.setattr(security, '_master', None)
    from clearline.app import create_app

    return create_app()


@pytest.fixture(scope='module')
def monkeypatch_module():
    mp = pytest.MonkeyPatch()
    yield mp
    mp.undo()


def sign_in(app, email, password):
    c = TestClient(app)
    r = c.post('/api/login', json={'email': email, 'password': password}, headers=H)
    assert r.status_code == 200, r.text
    body = r.json()
    secret = body.get('secret')
    if body['step'] == 'enroll':
        assert body['qr'].startswith('<?xml') or '<svg' in body['qr']
        SECRETS[email] = secret
    code = pyotp.TOTP(SECRETS[email]).now()
    r = c.post('/api/login/code', json={'pending': body['pending'], 'code': code}, headers=H)
    assert r.status_code == 200, r.text
    return c


SECRETS: dict[str, str] = {}
STATE: dict = {}


def test_setup_and_admin_login(app):
    c = TestClient(app)
    assert c.get('/api/setup').json()['needs_setup'] is True
    assert c.post('/api/setup', json={'name': 'Admin', 'email': 'admin@clearline.test', 'password': 'short'}, headers=H).status_code == 400
    assert c.post('/api/setup', json={'name': 'Admin', 'email': 'admin@clearline.test', 'password': 'Admin-pass-2026'}, headers=H).status_code == 200
    assert c.post('/api/setup', json={'name': 'X', 'email': 'x@x.test', 'password': 'Admin-pass-2026'}, headers=H).status_code == 409
    assert c.post('/api/login', json={'email': 'admin@clearline.test', 'password': 'wrong'}, headers=H).status_code == 401
    assert c.get('/api/me').status_code == 401
    STATE['admin'] = sign_in(app, 'admin@clearline.test', 'Admin-pass-2026')
    assert STATE['admin'].get('/api/me').json()['role'] == 'admin'


def test_header_required_for_changes(app):
    r = STATE['admin'].post('/api/users', json={'name': 'A', 'email': 'a@a.test', 'role': 'analyst'})
    assert r.status_code == 403


def test_people_and_clients(app):
    a = STATE['admin']
    people = {}
    for name, role in [('Sara', 'analyst'), ('Omar', 'analyst'), ('Mona', 'reviewer')]:
        r = a.post('/api/users', json={'name': name, 'email': f'{name.lower()}@clearline.test', 'role': role}, headers=H)
        assert r.status_code == 200, r.text
        people[name] = r.json()
    STATE['people'] = people
    clients = {}
    for slug, cur, analyst in [('nile-home-furnishings', 'EGP', 'Sara'), ('lotus-beauty', 'AED', 'Sara'), ('delta-medical-supplies', 'EGP', 'Omar')]:
        r = a.post('/api/clients', json={'name': slug.replace('-', ' ').title(), 'sector': 'retail', 'currency': cur, 'analyst_id': people[analyst]['id'], 'reviewer_id': people['Mona']['id']}, headers=H)
        assert r.status_code == 200, r.text
        clients[slug] = r.json()['id']
    STATE['clients'] = clients
    for name in ('Sara', 'Omar', 'Mona'):
        STATE[name] = sign_in(app, f'{name.lower()}@clearline.test', people[name]['temp_password'])
    assert STATE['Sara'].get('/api/me').json()['must_change_password'] is True
    assert STATE['Sara'].post('/api/me/password', json={'current': people['Sara']['temp_password'], 'new': 'Sara-new-pass-1'}, headers=H).status_code == 200


def test_access_is_limited_to_assigned_clients(app):
    sara, omar, mona = STATE['Sara'], STATE['Omar'], STATE['Mona']
    cl = STATE['clients']
    assert {c['id'] for c in sara.get('/api/clients').json()} == {cl['nile-home-furnishings'], cl['lotus-beauty']}
    assert sara.get(f'/api/clients/{cl["delta-medical-supplies"]}').status_code == 403
    assert omar.get(f'/api/clients/{cl["nile-home-furnishings"]}/files').status_code == 403
    assert len(mona.get('/api/clients').json()) == 3
    assert sara.post('/api/users', json={'name': 'Z', 'email': 'z@z.test', 'role': 'admin'}, headers=H).status_code == 403
    assert sara.patch(f'/api/clients/{cl["nile-home-furnishings"]}', json={'analyst_id': STATE['people']['Omar']['id']}, headers=H).status_code == 403
    log = STATE['admin'].get('/api/activity').json()
    assert log['chain_intact'] is True
    assert any('Blocked' in e['what'] for e in log['entries'])


def test_upload_map_and_run_everything(app):
    cl = STATE['clients']
    owner = {'nile-home-furnishings': STATE['Sara'], 'lotus-beauty': STATE['Sara'], 'delta-medical-supplies': STATE['Omar']}
    summary = {}
    for slug, cid in cl.items():
        c = owner[slug]
        files = [('files', (p.name, p.read_bytes(), 'text/csv')) for p in sorted((SAMPLES / slug).glob('*.csv'))]
        r = c.post(f'/api/clients/{cid}/files', files=files, headers=H)
        assert r.status_code == 200, r.text
        for f in r.json():
            assert f['kind'], f['name']
            detail = c.get(f'/api/files/{f["id"]}').json()
            assert detail['preview']
            r2 = c.put(f'/api/files/{f["id"]}/mapping', json={'kind': f['kind'], 'mapping': f['mapping'], 'options': {}}, headers=H)
            assert r2.status_code == 200, (f['name'], r2.text)
            assert r2.json()['ready'] and r2.json()['rows'] > 0
        svcs = c.get(f'/api/clients/{cid}/services').json()['services']
        assert len(svcs) == 32
        done = []
        for s in svcs:
            if s['level'] == 'none':
                continue
            t0 = time.time()
            r = c.post(f'/api/clients/{cid}/runs', json={'service_id': s['id'], 'params': s['params']}, headers=H)
            assert r.status_code == 200, r.text
            out = r.json()
            assert 'trace' not in out
            if 'result' in out:
                assert out['result']['kpis'] or out['result']['charts'], s['id']
                done.append((s['id'], out['id'], round(time.time() - t0, 1)))
        summary[slug] = done
        assert len(done) >= 20, (slug, len(done))
    STATE['runs'] = summary
    print({k: len(v) for k, v in summary.items()})


def test_report_review_and_export(app):
    sara, mona, admin = STATE['Sara'], STATE['Mona'], STATE['admin']
    cid = STATE['clients']['nile-home-furnishings']
    for sid, run_id, _ in STATE['runs']['nile-home-furnishings'][:6]:
        assert sara.post(f'/api/clients/{cid}/report/sections', json={'run_id': run_id}, headers=H).status_code == 200
    rep = sara.get(f'/api/clients/{cid}/report').json()
    assert len(rep['sections']) == 6
    first = rep['sections'][0]['id']
    sara.patch(f'/api/sections/{first}', json={'note': 'Sales are seasonal around Ramadan.'}, headers=H)
    sara.put(f'/api/clients/{cid}/report/summary', json={'text': 'Three things to fix this quarter. Sales grew 912.4% last year.'}, headers=H)
    rep = sara.get(f'/api/clients/{cid}/report').json()
    chk = rep['checks']
    assert chk['blocking'] == 0 and chk['passed'] > 20, chk
    assert any(i['id'] == 'summary-numbers' and '912.4%' in i['detail'] for i in chk['items'])  # a made-up figure is caught
    sara.put(f'/api/clients/{cid}/report/summary', json={'text': 'Three things to fix this quarter.'}, headers=H)
    assert mona.post(f'/api/clients/{cid}/report/review', json={'approve': True}, headers=H).status_code == 400  # not submitted
    assert sara.post(f'/api/clients/{cid}/report/submit', headers=H).json()['status'] == 'In review'
    assert sara.post(f'/api/clients/{cid}/report/review', json={'approve': True}, headers=H).status_code == 403
    assert mona.post(f'/api/clients/{cid}/report/review', json={'approve': False, 'comment': 'Explain the forecast range.'}, headers=H).json()['status'] == 'Changes requested'
    sara.post(f'/api/clients/{cid}/report/submit', headers=H)
    assert mona.post(f'/api/clients/{cid}/report/review', json={'approve': True}, headers=H).json()['status'] == 'Approved'
    docx = sara.get(f'/api/clients/{cid}/report.docx')
    assert docx.status_code == 200
    text = '\n'.join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert 'Sales are seasonal around Ramadan.' in text and 'Three things to fix' in text
    xlsx = sara.get(f'/api/clients/{cid}/report.xlsx')
    assert len(load_workbook(io.BytesIO(xlsx.content)).sheetnames) == 7
    out = sara.post(f'/api/clients/{cid}/report/deliver', headers=H).json()
    assert out['status'] == 'Delivered'
    out_dir = os.environ.get('CLEARLINE_REPORT_OUT')  # set it to keep the sample report files
    if out_dir:
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        (Path(out_dir) / 'sample_report.docx').write_bytes(docx.content)
        (Path(out_dir) / 'sample_report.xlsx').write_bytes(xlsx.content)


def test_failed_check_blocks_until_explained(app, monkeypatch):
    from clearline.engine import checks

    sara, admin = STATE['Sara'], STATE['admin']
    cid = STATE['clients']['lotus-beauty']
    again = sara.post(f'/api/clients/{cid}/runs', json={'service_id': 'business-snapshot', 'params': {}}, headers=H).json()
    assert next(k for k in again['result']['checks'] if k['id'] == 'stability')['detail'].startswith('Re-run on the same files')
    monkeypatch.setattr(checks, 'integrity', lambda res: checks.check('integrity', 'Result is complete', checks.FAIL, 'Planted failure.'))
    bad = sara.post(f'/api/clients/{cid}/runs', json={'service_id': 'trends-and-seasonality', 'params': {}}, headers=H).json()
    rep = sara.post(f'/api/clients/{cid}/report/sections', json={'run_id': bad['id']}, headers=H).json()
    sec = next(s for s in rep['sections'] if s['run']['id'] == bad['id'])
    assert rep['checks']['blocking'] == 1
    r = sara.post(f'/api/clients/{cid}/report/submit', headers=H)
    assert r.status_code == 400 and 'Planted failure' in r.json()['detail']
    assert sara.get(f'/api/clients/{cid}/report.docx').status_code == 400
    assert sara.post(f'/api/sections/{sec["id"]}/acks', json={'check_id': 'integrity', 'reason': 'fine'}, headers=H).status_code == 400
    assert sara.post(f'/api/sections/{sec["id"]}/acks', json={'check_id': 'plausible', 'reason': 'This one did not fail at all.'}, headers=H).status_code == 400
    rep = sara.post(f'/api/sections/{sec["id"]}/acks', json={'check_id': 'integrity', 'reason': 'Checked by hand against the client file.'}, headers=H).json()
    assert rep['checks']['blocking'] == 0 and rep['checks']['accepted'] == 1
    assert next(i for i in rep['checks']['items'] if i['id'] == 'integrity')['ack']['by'] == 'Sara'
    assert sara.post(f'/api/clients/{cid}/report/submit', headers=H).json()['status'] == 'In review'
    assert 'Accepted a failed check' in str(admin.get('/api/activity').json())


def test_delete_shreds_everything(app):
    from clearline import config

    admin = STATE['admin']
    cid = STATE['clients']['delta-medical-supplies']
    before = len(list(config.FILES_DIR.iterdir()))
    assert admin.delete(f'/api/clients/{cid}', headers=H).status_code == 200
    assert len(list(config.FILES_DIR.iterdir())) < before
    assert STATE['Omar'].get(f'/api/clients/{cid}').status_code == 404
    assert all(c['id'] != cid for c in admin.get('/api/clients').json())
    assert admin.get('/api/activity').json()['chain_intact'] is True


def test_lockout(app):
    c = TestClient(app)
    for _ in range(5):
        c.post('/api/login', json={'email': 'omar@clearline.test', 'password': 'nope'}, headers=H)
    r = c.post('/api/login', json={'email': 'omar@clearline.test', 'password': STATE['people']['Omar']['temp_password']}, headers=H)
    assert r.status_code == 429
